# -*- coding: utf-8 -*-
import base64
import csv
import io
import json
import secrets
import threading
import time
import zipfile
from .common import *

MAX_BYTES=5*1024*1024
MAX_ROWS=10000
FIELDS=['name','type','country','product','website','email','contact_name','source','source_date','facts','hypotheses','unknowns','signals','stage','blocker','priority']
def safe_csv(value):
    text=str(value)
    return "'"+text if text.lstrip().startswith(('=','+','-','@')) else text
class Importer:
    def __init__(self,store,demo):self.store=store;self.demo=demo;self.jobs={};self.lock=threading.RLock()
    def get(self,jid):
        job=self.jobs.get(jid)
        if not job or time.time()-job['created']>3600:raise AppError('导入任务不存在或已过期，请重新上传',404,'expired_import')
        return job
    def discard(self,jid):
        with self.lock:
            job=self.get(jid)
            if not job['committed']:self.jobs.pop(jid)
            return {'discarded':not bool(job['committed'])}
    def upload(self,data):
        filename=field(data,'filename',200,True);encoded=field(data,'content_base64',MAX_BYTES*2,True)
        try:raw=base64.b64decode(encoded,validate=True)
        except Exception:raise AppError('文件内容编码无效')
        if len(raw)>MAX_BYTES:raise AppError('文件最大5MB')
        rows=[];headers=[];row_numbers=[]
        if filename.lower().endswith('.csv'):
            try:text=raw.decode('utf-8-sig')
            except UnicodeDecodeError:
                try:text=raw.decode('gb18030')
                except UnicodeDecodeError:raise AppError('CSV请使用UTF-8或GB18030编码')
            try:
                sample=text[:4096]
                try:dialect=csv.Sniffer().sniff(sample,delimiters=',;\t')
                except csv.Error:dialect=csv.excel
                reader=csv.reader(io.StringIO(text),dialect,strict=True);headers=next(reader,[]);previous_line=reader.line_num
                for row in reader:
                    original_line=previous_line+1;previous_line=reader.line_num
                    if len(rows)>=MAX_ROWS:raise AppError('一次最多10000行，请分文件导入')
                    if any(str(v).strip() for v in row):rows.append(row);row_numbers.append(original_line)
            except csv.Error:raise AppError('CSV格式无法解析')
        elif filename.lower().endswith('.xlsx'):
            try:
                from openpyxl import load_workbook
                with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                    infos=archive.infolist()
                    if len(infos)>2000 or sum(x.file_size for x in infos)>50*1024*1024:raise AppError('XLSX解压内容超过安全限制')
                    if any('vbaProject' in x.filename for x in infos):raise AppError('不接受带宏工作簿')
                book=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
                sheet=book.worksheets[0];iterator=sheet.iter_rows()
                first=next(iterator,[])
                if any(c.data_type=='f' for c in first):raise AppError('表头不接受公式')
                headers=[str(c.value or '') for c in first]
                for physical_row,cells in enumerate(iterator,2):
                    if physical_row>MAX_ROWS+1:raise AppError('工作表超10000行（包括空白行），请清理后导入')
                    if len(rows)>=MAX_ROWS:raise AppError('一次最多10000行')
                    # Formula cells are rejected instead of executed or evaluated.
                    row=[('__FORMULA_NOT_ALLOWED__' if c.data_type=='f' else str(c.value.isoformat()[:10] if hasattr(c.value,'isoformat') else c.value or '')) for c in cells]
                    if any(v.strip() for v in row):rows.append(row);row_numbers.append(physical_row)
                book.close()
            except ImportError:raise AppError('缺少openpyxl，请按README安装项目依赖',500,'missing_dependency')
            except AppError:raise
            except Exception:raise AppError('XLSX无法解析，请使用普通.xlsx工作簿')
        else:raise AppError('仅支持.csv与.xlsx，不接受.xls或.xlsm')
        headers=[str(x).strip() for x in headers]
        if not headers or len(headers)>50 or any(not h or len(h)>100 for h in headers) or len(set(headers))!=len(headers):raise AppError('表头须唯一且非空，最多50列，每列名称最多100字符')
        if not rows:raise AppError('文件没有数据行')
        if any(len(row)>50 or any(len(str(v))>12000 for v in row) for row in rows):raise AppError('文件列数或单元格长度超限')
        job={'id':secrets.token_hex(12),'created':time.time(),'filename':filename,'headers':headers,'raw':rows,'row_numbers':row_numbers,'preview':None,'committed':None}
        with self.lock:
            expired=[k for k,v in self.jobs.items() if time.time()-v['created']>3600]
            for k in expired:self.jobs.pop(k)
            if len(self.jobs)>=5:raise AppError('最多保留5个导入任务，请完成或等待旧任务过期')
            self.jobs[job['id']]=job
        return {'job_id':job['id'],'headers':headers,'preview_rows':rows[:5],'total_rows':len(rows),'fields':FIELDS}
    def preview(self,data):
        with self.lock:
            job=self.get(field(data,'job_id',32,True))
            if job['committed']:raise AppError('已确认任务不可修改映射',409,'already_committed')
            mapping=data.get('mapping');defaults=data.get('defaults',{});action=data.get('duplicate_action','skip')
            if not isinstance(mapping,dict) or not isinstance(defaults,dict) or action not in ['skip','keep']:raise AppError('映射或重复处理参数无效')
            if not mapping.get('name'):raise AppError('必须映射公司名称')
            if any(k not in FIELDS or v not in job['headers'] for k,v in mapping.items() if v):raise AppError('映射字段无效')
            with self.store.db() as db:existing=set(db.execute('SELECT name_key,country_key FROM customers'))
            seen=set();preview=[]
            for raw_index,row in enumerate(job['raw']):
                index=job['row_numbers'][raw_index]
                source=dict(zip(job['headers'],row));c={key:source.get(col,'').strip() for key,col in mapping.items() if col};issues=[]
                if any('__FORMULA_NOT_ALLOWED__' in v for v in row):issues.append('公式单元格不接受，请替换为核实后的纯文本')
                if len(row)!=len(job['headers']):issues.append('列数与表头不一致')
                c['type']=c.get('type') or defaults.get('type') or TYPES[0];c['stage']=c.get('stage') or defaults.get('stage') or STAGES[0];c.setdefault('country',defaults.get('country',''))
                c.setdefault('source','用户导入文件：'+job['filename']+'；来源待核实')
                if isinstance(c.get('priority'),str):
                    value=c['priority'].strip().lower()
                    if value not in ['','0','1','true','false','是','否']:issues.append('重点标记须是/否或true/false')
                    c['priority']=value in ['1','true','是']
                try:c=customer_data(c)
                except AppError as e:issues.append(str(e))
                identity=(norm(c.get('name','')),norm(c.get('country','')))
                duplicate=bool(identity[0]) and (identity in existing or identity in seen)
                if not issues:seen.add(identity)
                status='invalid' if issues else 'duplicate' if duplicate else 'valid'
                preview.append({'row':index,'raw_index':raw_index,'status':status,'issues':issues or (['同名及国家匹配：疑似重复，不自动覆盖'] if duplicate else []),'customer':c})
            job['preview']=preview;job['duplicate_action']=action
            return {'job_id':job['id'],'valid_count':sum(r['status']=='valid' for r in preview),'invalid_count':sum(r['status']=='invalid' for r in preview),'duplicate_count':sum(r['status']=='duplicate' for r in preview),'rows':preview[:100],'total_rows':len(preview),'preview_limit':100}
    def confirm(self,data):
        if data.get('consent') is not True:raise AppError('请明确确认导入')
        with self.lock:
            jid=field(data,'job_id',32,True)
            with self.store.db() as db:receipt=db.execute('SELECT data FROM import_receipts WHERE id=?',(jid,)).fetchone()
            if receipt:return dict(json.loads(receipt[0]),already_committed=True)
            job=self.get(jid)
            if job['committed']:return dict(job['committed'],already_committed=True)
            if job['preview'] is None:raise AppError('请先映射并预览校验')
            inserted=skipped=failed=0
            with self.store.db() as db:
                db.execute('BEGIN IMMEDIATE')
                receipt=db.execute('SELECT data FROM import_receipts WHERE id=?',(jid,)).fetchone()
                if receipt:return dict(json.loads(receipt[0]),already_committed=True)
                existing=set(db.execute('SELECT name_key,country_key FROM customers'))
                for row in job['preview']:
                    if row['status']=='invalid':failed+=1;continue
                    c=row['customer'];identity=(norm(c['name']),norm(c['country']))
                    # Recheck after preview: concurrent inserts must not silently create duplicates.
                    if identity in existing and job['duplicate_action']=='skip':skipped+=1;continue
                    c=dict(c,id=secrets.token_hex(8),demo=self.demo,version=1,updated_at=now());self.store.insert(db,c);existing.add(identity);inserted+=1
                result={'inserted':inserted,'skipped':skipped,'failed':failed,'already_committed':False}
                db.execute('INSERT INTO import_receipts VALUES(?,?)',(jid,dump(result)))
            job['committed']=result;return result
    def errors(self,jid):
        with self.lock:
            job=self.get(jid)
            if job['preview'] is None:raise AppError('请先进行校验预览')
            out=io.StringIO();writer=csv.writer(out);writer.writerow(['原始行号','状态','问题']+job['headers'])
            for row in job['preview']:
                if row['status']=='invalid':writer.writerow([row['row'],row['status'],'；'.join(row['issues'])]+[safe_csv(x) for x in job['raw'][row['raw_index']]])
            return ('\ufeff'+out.getvalue()).encode('utf-8')
