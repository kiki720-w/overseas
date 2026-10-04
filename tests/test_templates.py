# -*- coding: utf-8 -*-
import test_server
import unittest
import csv
import io
import urllib.request
from openpyxl import load_workbook

class TemplateTests(unittest.TestCase):
    setUp=test_server.WorkspaceTests.setUp
    tearDown=test_server.WorkspaceTests.tearDown
    call=test_server.WorkspaceTests.call
    def test_template_downloads_and_structure(self):
        with urllib.request.urlopen(self.base+'/templates/customers.xlsx') as r:
            self.assertIn('attachment',r.headers.get('Content-Disposition',''))
            wb=load_workbook(io.BytesIO(r.read()),read_only=True)
            self.assertEqual(wb.worksheets[0]['A1'].value,'公司名称')
            self.assertEqual(wb.worksheets[0].max_row,1)
            wb.close()
        for route,count in [('/templates/customers.csv',1),('/templates/demo.csv',3)]:
            with urllib.request.urlopen(self.base+route) as r:
                rows=list(csv.reader(io.StringIO(r.read().decode('utf-8-sig'))))
                self.assertEqual(len(rows),count)
                self.assertEqual(rows[0][0],'公司名称')
    def test_template_file_paths_are_not_open(self):
        import urllib.error
        with self.assertRaises(urllib.error.HTTPError) as ctx:self.call('/templates/../server.py')
        self.assertEqual(ctx.exception.code,404)
